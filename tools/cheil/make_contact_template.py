# Cheil fork: contact report authoring.
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
root=Path(os.environ['CHEIL_TEMPLATE_OUTPUT']).parent
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='맑은 고딕';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕');s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.15
d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(14);d.styles['Heading 2'].font.size=Pt(11)
def para(t,style=None):return d.add_paragraph(t,style)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
 for ri,row in enumerate(t.rows):
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF2' if ri==0 else 'FFFFFF');pr.append(shade)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    for run in p.runs:run.bold=ri==0
  if ri==0:
   rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:tblHeader'))
 return t
para('컨택리포트','Title');para('CONTACT REPORT')
table(['항목','내용'],[
 ['캠페인명','[캠페인명]'],['광고주 · 브랜드·제품','[광고주] · [브랜드·제품]'],['회의명','[회의명]'],['Date / Place','[회의 일시]  /  [회의 장소]'],['참석자','광고주  [광고주 참석자]\n우리 측  [우리 측 참석자]'],['작성일 · 팀명 · 작성자','[작성일] · [팀명] · [작성자]']], [4.2,12.8])
para('주요 결정과 확인 사항','Heading 1');para('[회의 핵심 결정]');para('[결정되지 않은 사항]')
para('상세 피드백','Heading 1')
for i in [1,2]:
 para(f'안건 {i}  [안건명 {i}]','Heading 2')
 para(f'광고주 피드백  [광고주 피드백 {i}]')
 para(f'합의 및 수정 요청  [합의와 수정 요청 {i}]')
para('AE 의견','Heading 1');para('[AE 의견과 대응 방향]')
para('NEXT','Heading 1')
table(['할 일','담당','기한','확인 또는 전달물'],[
 ['[후속 할 일 1]','[후속 담당 1]','[후속 기한 1]','[후속 전달물 1]'],
 ['[후속 할 일 2]','[후속 담당 2]','[후속 기한 2]','[후속 전달물 2]'],
 ['[후속 할 일 3]','[후속 담당 3]','[후속 기한 3]','[후속 전달물 3]']],[5.5,3.3,3.2,5.0])
para('추가 확인 사항','Heading 1');para('[추가 확인 항목과 확인할 사람]')
p=para('- END -');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
sec.footer.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');sec.footer.paragraphs[0]._p.append(fld)
for tree in [d.styles.element,d.element]:
 for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.save(Path(os.environ['CHEIL_TEMPLATE_OUTPUT'])/'컨택리포트.docx')
print('contact report saved')
